from __future__ import annotations
from datetime import datetime
import math, re
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from core.paths import OUTPUT_DIR
from agent.risk_brief import baseline_brief, scenario_brief

CJK="Noto Sans CJK SC"

def _missing(v):
    return v is None or (isinstance(v,float) and (math.isnan(v) or math.isinf(v)))

def _fmt(v,n=4):
    if _missing(v): return "不可计算"
    try: return f"{float(v):.{n}f}"
    except Exception: return "不可计算"

def _pct(v):
    if _missing(v): return "不可计算"
    try: return f"{float(v)*100:.1f}%"
    except Exception: return "不可计算"

def _clean_md(s:str)->str:
    s=re.sub(r"\*\*(.*?)\*\*",r"\1",s); return s.replace("`","").strip()

def _add_brief(doc,text):
    for raw in (text or "").splitlines():
        line=raw.strip()
        if not line: continue
        if line.startswith("### "): doc.add_heading(_clean_md(line[4:]),level=2)
        elif line.startswith("- "): doc.add_paragraph(_clean_md(line[2:]),style="List Bullet")
        else: doc.add_paragraph(_clean_md(line))

def _force_cjk(doc):
    for st in doc.styles:
        try:
            st.font.name=CJK; st._element.rPr.rFonts.set(qn("w:eastAsia"),CJK)
        except Exception: pass
    for para in doc.paragraphs:
        for run in para.runs:
            run.font.name=CJK; run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),CJK)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.name=CJK; run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"),CJK)

def export_risk_report(baseline_result:dict, scenario_result:dict|None, validation:dict|None, snapshot_id:str="",
                       baseline_explanation:str="", scenario_explanation:str="")->str:
    doc=Document(); doc.styles["Normal"].font.name=CJK; doc.styles["Normal"].font.size=Pt(10.5)
    doc.add_heading("WaterPulse 上游水风险分析报告",0)
    doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    if snapshot_id: doc.add_paragraph(f"本次运行编号：{snapshot_id}")

    doc.add_heading("1. 数据状态",1)
    if validation:
        doc.add_paragraph("资料检查状态："+str(validation.get("status") or "Unknown"))
        for k,cn in [("issues","需要确认"),("warnings","提示"),("data_gaps","数据缺口")]:
            vals=validation.get(k,[])
            if vals: doc.add_paragraph(cn+"："+"；".join(map(str,vals)))
    ident=(baseline_result or {}).get("data_identity") or {}
    if ident: doc.add_paragraph("数据身份："+str(ident.get("label") or "未标记")+"。"+str(ident.get("note") or ""))
    proxies=(baseline_result or {}).get("proxy_fields") or []; assumptions=(baseline_result or {}).get("assumptions") or []
    if proxies: doc.add_paragraph(f"代理数据：{len(proxies)} 项。")
    if assumptions: doc.add_paragraph(f"假设参数：{len(assumptions)} 项。")

    doc.add_heading("2. 当前风险结果",1)
    brief=baseline_explanation.strip() or baseline_brief(baseline_result)
    _add_brief(doc,brief)
    if baseline_result and baseline_result.get("data") is not None:
        s=baseline_result["data"]["summary"]; nodes=baseline_result["data"]["nodes"]
        if isinstance(s,list): import pandas as pd; s=pd.DataFrame(s)
        if isinstance(nodes,list): import pandas as pd; nodes=pd.DataFrame(nodes)
        table=doc.add_table(rows=1,cols=9); table.style="Table Grid"
        heads=["企业","原材料","综合风险","采购覆盖","有效数据覆盖","未识别份额","长期缺水贡献","历史干旱贡献","季节波动贡献"]
        for i,x in enumerate(heads): table.rows[0].cells[i].text=x
        for _,r in s.iterrows():
            total=r.get("PRWI"); ws=(r.get("path_contrib_WS")/total if total not in (None,0) else None); dr=(r.get("path_contrib_DR")/total if total not in (None,0) else None); sv=(r.get("path_contrib_SV")/total if total not in (None,0) else None)
            vals=[r.get("enterprise"),r.get("material"),_fmt(total),_pct(r.get("Coverage")),_pct(r.get("scored_coverage")),_pct(r.get("unknown_share_U")),_pct(ws),_pct(dr),_pct(sv)]
            cells=table.add_row().cells
            for i,v in enumerate(vals): cells[i].text=str(v if v not in (None,"") else "不可计算")
            if hasattr(nodes,"empty") and not nodes.empty:
                gn=nodes[nodes.material==r.get("material")] if "material" in nodes.columns else nodes
                if not gn.empty and "C" in gn.columns:
                    top=gn.sort_values("C",ascending=False).iloc[0]
                    doc.add_paragraph(f"{r.get('material')} 最大贡献供应地：{top.get('node_name')}（{top.get('node_id')}），贡献占比 {_pct(top.get('contribution_share'))}。")

    doc.add_heading("3. 不同情况下的变化",1)
    _add_brief(doc,(scenario_explanation.strip() if scenario_result else "") or (scenario_brief(scenario_result) if scenario_result else "本次未进行额外压力情况比较。"))

    doc.add_heading("4. 来源与可追溯信息",1)
    details=(baseline_result or {}).get("source_details") or []
    if details:
        table=doc.add_table(rows=1,cols=6); table.style="Table Grid"
        for i,x in enumerate(["对象","项目来源标识","参考期","空间匹配","数据身份","置信度"]): table.rows[0].cells[i].text=x
        for d in details[:30]:
            obj=(d.get("node_name") or d.get("material") or d.get("node_id") or "数据项")
            vals=[obj,d.get("source_ref") or "待补",d.get("reference_period") or "待补",d.get("pfaf_id") or "待补",d.get("data_type") or "待补",d.get("confidence") or "待补"]
            cells=table.add_row().cells
            for i,v in enumerate(vals): cells[i].text=str(v)
        doc.add_paragraph("说明：当前项目文件中尚未完整登记原始来源机构和公开链接，因此报告只展示项目已注册来源标识、参考期、地点匹配和数据身份；原始数据库链接需由数据负责人补齐，系统不会自行杜撰。")
    else:
        doc.add_paragraph("当前结果没有可展示的结构化来源明细。")

    doc.add_heading("5. 使用边界与建议",1)
    doc.add_paragraph("当前 MVP 的自动地点匹配仅覆盖项目已登记的供应节点。未登记地区不会被静默替换为其他地区；如需新增地区，应先补充并核验数据。")
    doc.add_paragraph("综合风险指数用于相对风险筛查，不是损失概率或财务预测。实际采购调整、供应商替换和投资决策仍需人工确认。")
    doc.add_heading("6. 版本",1)
    doc.add_paragraph(f"风险计算引擎：{(baseline_result or {}).get('engine_version') or '待记录'}；数据版本：{(baseline_result or {}).get('data_version') or '待记录'}。")
    _force_cjk(doc)
    out=OUTPUT_DIR/f"WaterPulse_Risk_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"; doc.save(out); return str(out)
